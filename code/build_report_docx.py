"""Build the Word deliverable from the completed replication report content."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "replication-lab" / "replication-lab.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10.5)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    styles["Normal"].paragraph_format.line_spacing = 1.08
    for name, size in (("Title", 22), ("Heading 1", 15), ("Heading 2", 12)):
        style = styles[name]
        style.font.name = "Aptos Display" if name == "Title" else "Aptos"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)
    # Word's built-in Title style may carry a decorative bottom rule. Remove it
    # so the report title remains plain black and consistent with the lab template.
    title_ppr = styles["Title"]._element.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)


def add_para(doc, text: str, bold_lead: str | None = None) -> None:
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        p.add_run(bold_lead).bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)


def add_metadata(doc, label: str, value: str) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    p.add_run(label + ": ").bold = True
    p.add_run(value)


def add_results_table(doc: Document) -> None:
    headers = ["Agent", "Task A\nbefore", "Task B\nbefore", "Task A\nfinal", "Task B\nfinal", "Task A\ndrop", "Epoch to\nTask B 80%"]
    rows = [
        ["Replay stability", "0.993", "0.499", "0.898", "0.618", "0.095", "Not reached"],
        ["No replay", "0.993", "0.498", "0.507", "0.991", "0.486", "2"],
        ["Fresh Task B benchmark", "--", "--", "0.502", "0.988", "--", "1"],
    ]
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        set_cell_shading(cell, "2F5597")
        set_cell_borders(cell)
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.size = Pt(8.5)
    for row_i, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cell = cells[i]
            cell.text = value
            set_cell_borders(cell)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_i % 2 == 1:
                set_cell_shading(cell, "F2F5FA")
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8.5)


def add_figure(doc: Document, path: Path, caption: str) -> None:
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(6)
    p.add_run().add_picture(str(path), width=Inches(6.15))
    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.paragraph_format.space_after = Pt(8)
    for run in cap.runs:
        run.italic = True
        run.font.size = Pt(9)


def add_bullets(doc, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.add_run(item)


def main() -> None:
    doc = Document()
    style_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("Replication Laboratory 2 Can the Model Keep Learning")
    title_ppr = title._p.get_or_add_pPr()
    title_borders = title_ppr.find(qn("w:pBdr"))
    if title_borders is not None:
        title_ppr.remove(title_borders)
    subtitle = doc.add_paragraph("ANLY 735 Research Seminar in Predictive AI")
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.runs[0].italic = True
    subtitle.runs[0].font.color.rgb = RGBColor(80, 80, 80)

    add_metadata(doc, "Student", "Huaqing Fang")
    add_metadata(doc, "Article", "Plasticity Loss in Deep Reinforcement Learning A Survey")
    add_metadata(doc, "Authors", "Timo Klein, Lukas Miklautz, Kevin Sidak, Claudia Plant, and Sebastian Tschiatschek")
    add_metadata(doc, "Replication type", "Proxy replication")
    add_metadata(doc, "Environment", "Python; PyTorch 2.13.0, NumPy 2.5.2, pandas 3.0.5, matplotlib 3.11.1")
    add_metadata(doc, "Random seed", "20250927")
    add_metadata(doc, "Repository", "https://github.com/hf2431/anly735-lab02-hf2431")

    doc.add_heading("1 Research Claim", level=1)
    doc.add_heading("What result are you attempting to reproduce", level=2)
    add_para(doc, "The anchor paper is a survey, so it does not provide one original training run that can be reproduced exactly. I therefore focus on its broader stability–plasticity claim: when a learning problem changes, an already-trained neural network can retain some capability from the earlier condition while becoming less effective at learning what comes next. This is important because a lower score immediately after a change is not enough to diagnose plasticity loss. The more informative evidence is the post-change trajectory: does the model improve, how quickly does it improve, and what previous capability is retained while it adapts? My small experiment treats these as separate outcomes rather than reducing the question to one final accuracy. The target is a narrow proxy for the phenomenon synthesized by Klein et al. (2024).")

    doc.add_heading("2 Replication Strategy", level=1)
    doc.add_heading("How will you investigate the claim", level=2)
    add_para(doc, "This is a proxy replication. The original survey synthesizes findings from many studies, but it does not supply a single dataset, model checkpoint, and learning trajectory for a direct rerun. I constructed a controlled two-task experiment that makes the relevant change visible. Two identical multilayer perceptrons first learn Task A. After the task changes to Task B, one persistent agent receives Task B examples plus a small replay stream from Task A, while the other receives only Task B examples. A freshly initialized Task B model provides a simple learning-rate benchmark. I compare Task A retention, Task B accuracy, and the number of post-change epochs needed to reach 80% Task B accuracy. This does not test every mechanism in the survey; it tests whether a stability–plasticity tradeoff can be observed transparently in a reproducible toy setting.")

    doc.add_heading("3 Data", level=1)
    add_para(doc, "The data are synthetic and generated by code/lab02_analysis.py, so no restricted or personally identifiable data are involved. Each task contains 2,000 training observations and a separate 2,000-observation evaluation set. Each observation has two independent standard-normal features. In Task A, the binary target is whether feature 1 is positive. In Task B, the target is whether feature 2 is positive. The task switch therefore changes the predictive rule while keeping the sample format and difficulty comparable. The unit of analysis is one simulated observation and the outcome is a binary class label. There are no missing values and no external preprocessing steps. The synthetic construction differs from the reinforcement-learning settings reviewed in the anchor paper, but it gives a clean way to separate retention of the old rule from learning of the new rule.")

    doc.add_heading("4 Method", level=1)
    add_para(doc, "The model is a small multilayer perceptron with two input features, one hidden layer of 32 ReLU units, and a single sigmoid output. Both persistent agents begin from the same seeded initialization and train on Task A for 10 epochs using stochastic gradient descent with learning rate 0.10 and batch size 128. At the condition change, the replay stability agent trains on Task B while also replaying Task A examples with replay weight 2.0. The no-replay agent trains only on Task B. A fresh Task B benchmark is initialized independently and trained only on Task B. All agents are evaluated after each post-change epoch on held-out Task A and Task B data.")
    add_para(doc, "The main measures are Task A accuracy before and after the change, the Task A retention drop, the final Task B accuracy, and the first post-change epoch at which Task B accuracy reaches 80%. The fixed seed is 20250927. The analysis can be rerun from the repository root with python code/lab02_analysis.py; it writes the trajectory and summary CSV files and saves both figures in figures/.")

    doc.add_heading("5 Results", level=1)
    add_para(doc, "The generated results are summarized below. Accuracies are evaluated on held-out data. The retention drop is the Task A accuracy before the change minus the final Task A accuracy.")
    add_results_table(doc)
    add_figure(doc, ROOT / "figures" / "task_b_learning.png", "Figure 1. New learning after the condition change.")
    add_figure(doc, ROOT / "figures" / "task_a_retention.png", "Figure 2. Retention of the previously learned task.")
    add_para(doc, "The curves show the tradeoff more clearly than the final scores alone. The no-replay agent learns the new rule quickly, reaching 80% Task B accuracy by epoch 2 and ending at 0.991, but its Task A accuracy falls from 0.993 to 0.507. The replay agent protects more of the old capability, ending at 0.898 on Task A with only a 0.095 drop, but it improves slowly on Task B and never reaches the 80% threshold in the ten post-change epochs. The fresh benchmark learns Task B quickly, showing that the new task itself is not unusually difficult.")

    doc.add_heading("6 Compare", level=1)
    doc.add_heading("How closely did your results align with the original study", level=2)
    add_para(doc, "Because the anchor paper is a survey rather than one empirical experiment, an identical numerical comparison is not appropriate. The meaningful comparison is directional: does a controlled nonstationary learning problem make retention and new learning pull in different directions? My result is consistent with that broader methodological idea. The no-replay agent has high plasticity for the new rule but poor stability for the old rule. The replay agent has stronger stability but slower new learning. The fresh benchmark confirms that a newly initialized model can learn Task B rapidly, so the replay agent's slow improvement is associated with carrying forward the prior task and the imposed replay constraint, not simply with Task B being impossible.")
    add_para(doc, "This comparison also limits the claim. The experiment does not show that an ordinary neural network spontaneously loses plasticity in every setting, and it does not estimate a universal effect size. Replay was deliberately introduced as a stabilizing condition, and the tasks are simple synthetic classification rules. The evidence supports a stability–plasticity tradeoff in this controlled proxy, not a direct reproduction of all findings reviewed by Klein et al. (2024).")

    doc.add_heading("7 Reproducibility Check", level=1)
    add_bullets(doc, [
        "✓ Data source documented",
        "✓ Code runs from beginning to end",
        "✓ Computational environment identified",
        "✓ Software and package versions documented",
        "✓ Random seed documented",
        "✓ Key preprocessing steps documented",
        "✓ Evaluation procedure documented",
        "✓ Repository README provides reproduction instructions",
    ])
    add_para(doc, "The main challenge was scope rather than execution. Since the anchor paper is a survey, there was no single original dataset or script to download and rerun. I addressed that by making the proxy experiment explicit, fixing the random seed, keeping the two persistent agents' starting conditions identical, and writing the generated CSV files and figures to named repository locations. This makes the design easy to inspect, while also making clear that the numerical results should not be presented as a direct replication of a particular study in the survey.", bold_lead="The greatest reproducibility challenge was scope rather than execution. ")

    doc.add_heading("8 Extend", level=1)
    doc.add_heading("What did the replication reveal that you would investigate next", level=2)
    add_para(doc, "The next useful extension would be to repeat the same experiment over many random seeds and over a sequence of several task changes instead of only one switch. I would compare no replay, replay, and a plasticity-preserving method under the same sequence, then report the distribution of retention drops, recovery rates, and time to the 80% threshold. This would distinguish a stable pattern from one outcome of a particular initialization and would show whether the tradeoff accumulates after repeated changes. It adds knowledge because it tests the robustness and persistence of the observed learning behavior rather than merely adding a larger or more complicated network.")

    doc.add_heading("Replication Verdict", level=1)
    add_bullets(doc, [
        "☐ Reproduced — The central result was substantially reproduced.",
        "☒ Partially Reproduced — Some findings were reproduced, but meaningful differences remain.",
        "☐ Not Reproduced — The central result was not reproduced.",
        "☐ Inconclusive — The evidence was insufficient to support a directional conclusion.",
    ])
    add_para(doc, "I classify this focused proxy replication as Partially Reproduced. After the condition change, the replay agent retained most of Task A but learned Task B slowly, while the no-replay agent learned Task B quickly but lost Task A; this is evidence of the stability–plasticity tradeoff described by the broader literature. The result does not justify claiming that the entire Klein et al. survey, or spontaneous plasticity loss in general, has been reproduced.")

    doc.add_heading("References", level=1)
    add_para(doc, "Klein, Timo, Lukas Miklautz, Kevin Sidak, Claudia Plant, and Sebastian Tschiatschek. 2024. “Plasticity Loss in Deep Reinforcement Learning: A Survey.” arXiv preprint arXiv:2411.04832. https://arxiv.org/abs/2411.04832")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
